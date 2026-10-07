from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
import json, sys, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment
sys.path.insert(0, str(Path(__file__).resolve().parent))
import lc_data as D

SRC = str(ROOT / "inputs/torc") + "/"
OUT = ROOT / "catalogue/LaneChange_Scenario_Library_Foretellix_style_v0.xlsx"
ref = json.load(open(ROOT / "data/derived/ref.json"))
TC, FI = ref["tc"], ref["fi"]

# HIRE Highway Lane Change rows
wb = openpyxl.load_workbook(SRC + "2026_SDT_HIRE_from_HAZOP_v1.xlsx", read_only=True, data_only=True)
HIRE = []
for r in list(wb["HIRE"].iter_rows(values_only=True))[1:]:
    if r[2] == "Highway" and "Lane Change" in str(r[3]):
        HIRE.append(dict(id=r[0], odd=r[4], man=r[5], ctx=r[6], gw=r[7], mal=r[8], ev=r[9], hz=r[10], S=r[11], E=r[13], C=r[15], hazop=r[21]))

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
FAM_FILL = {"LC-NOM": "E2EFDA", "LC-TGT": "DDEBF7", "LC-CUR": "FFF2CC", "LC-RD": "FCE4D6", "LC-ABT": "EDE7F6", "LC-OOD": "F4B084"}
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

def style_header(ws, row, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = HDR_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = BORDER

def write_table(ws, start_row, headers, rows, widths, fam_col=None):
    for j, h in enumerate(headers, 1):
        ws.cell(row=start_row, column=j, value=h)
    style_header(ws, start_row, len(headers))
    for i, r in enumerate(rows, start_row + 1):
        for j, v in enumerate(r, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = Font(name=FONT, size=9)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = BORDER
        if fam_col is not None:
            fam = r[fam_col]
            if fam in FAM_FILL:
                for j in range(1, len(headers) + 1):
                    ws.cell(row=i, column=j).fill = PatternFill("solid", fgColor=FAM_FILL[fam])
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=start_row + 1, column=3)
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{start_row + len(rows)}"

def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = Font(name=FONT, bold=True, size=14, color="1F3864")
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(name=FONT, italic=True, size=9, color="595959")

def split(s):
    return [x.strip() for x in s.split(",") if x.strip()]

# ---------- derived per scenario ----------
for s in D.S:
    tcs = [t for t in split(s["tcs"]) if t in TC]
    missing = [t for t in split(s["tcs"]) if t not in TC]
    assert not missing, (s["id"], missing)
    fis = sorted({f for t in tcs for f in TC[t]["fi"]})
    s["fis"] = fis
    hz_tc = {h.strip() for t in tcs for h in (TC[t]["haz"] or "").split(",") if h.strip()}
    rows = [h for h in HIRE if h["ctx"] in s["ctx"]]
    s["hire_rows"] = rows
    hz_hire = {h["hz"] for h in rows}
    s["hz"] = sorted(hz_tc | hz_hire)
    s["mal"] = sorted({h["mal"] for h in rows})
    s["hfs"] = "; ".join(f"Highway x Lane Changes / Overtaking x {h}" for h in sorted(hz_hire))
    s["tc_txt"] = "\n".join(f"{t}: {TC[t]['tc']}" for t in tcs)
    s["fi_txt"] = "\n".join(f"{f}: {FI[f]['name']}" for f in fis)

# reverse map HIRE -> scenarios
for h in HIRE:
    h["sc"] = [s["id"] for s in D.S if h["ctx"] in s["ctx"]]

out = openpyxl.Workbook()

# ---------- Read Me ----------
ws = out.active
ws.title = "Read Me"
title(ws, "Lane Change Scenario Library, Foretellix style (v0, PROPOSED)",
      "SERYTI for Torc Robotics, DO-Crawl I-35 South. Behavior: Highway 'Lane Changes / Overtaking' (Torc Behavior Table: Safety, Traffic Compliance, Efficiency, Courtesy).")
lines = [
    ("PURPOSE", "First full Phase 1A deliverable of the Miro flow for one behavior: Behavior > Scenario decomposition (static + dynamic ODD elements) > SDT tasks > wanted behavior rules, linked to Phase 1B (unwanted / unsafe behaviors from HIRE) and to TC / FI. Built as a Foretellix-style scenario library so it can be executed in a scenario-based tool (OpenSCENARIO DSL)."),
    ("LEVELS", "Torc Miro 'Abstract Scenario' (ego maneuver only) = scenario FAMILY root here (LC-NOM-01).\nTorc Miro 'Concrete Scenario A / B' (ego + placed actors) = ABSTRACT scenarios in Foretellix / ISO 34501 terms (LC-NOM-02, LC-TGT-02): actors, phases, symbolic parameters.\nVariants of an actor (Miro: E holds SDT speed / E accelerates) = parameter values, swept at LOGICAL level (ranges, sheet 'Parameters') then CONCRETE level (sampled values, done in the tool)."),
    ("FORETELLIX MODEL", "Each abstract scenario carries what a Foretellix V-Suite scenario carries: actors with roles, named phases (serial / parallel), parameters with constraints, coverage items with buckets, KPIs and checkers, and mixable modifiers. Public structure reference: Foretellix evaluation V-Suite docs (e.g. sut_cut_out, vehicle_cut_in). The OSC2 file 'LaneChange_Library.osc' gives the same scenarios as pseudo OpenSCENARIO DSL (not compiled)."),
    ("SLOT MODEL", "Actor positions follow the Torc Miro grid: A lead origin lane, B lead target lane, C follower origin lane, D follower target lane, E / F lane beyond target, G other-side lane, S shoulder, O static object. See sheet 'Actor Slots'."),
    ("MIRO OBJECTIVES COVERED", "Wanted behavior specs and rules: column 'Wanted behavior (rule book)'.\nUnwanted / unsafe behaviors: 'HIRE Traceability' (61 Highway Lane Change HARA) + columns 'Hazards', '# HIRE rows'.\nExternal factors: 'Triggering conditions (TC-CRAWL)' + 'Modifiers'.\nInternal factors: 'Functional insufficiencies (FI)'.\nInternal factors along the task chain + KPIs on task I/O (solution independent): sheet 'Task Chain' (T1..T10). SOTIF team stops there; solution-dependent causes and KPIs are out of scope."),
    ("SHEETS", "Library Tree: families and scenarios at a glance.\nScenario Catalogue: one row per abstract scenario (main sheet).\nActor Slots, Parameters (logical ranges with ODD source), KPIs & Checkers, Task Chain, Modifiers.\nHIRE Traceability: every Lane Change HARA row and the scenarios that exercise it.\nScenario x TC x FI: exploded join.\nCoverage Stats: live counts (formulas).\nGaps & Open Points."),
    ("LINK RULES", "HIRE rows are linked to a scenario through their HAZOP 'Hazard Context' (the competency the scenario stresses). FIs are derived from the scenario's TCs through Full_Mapping_TC_FI_Hazard_Arch.xlsx. Hazards = union of linked HIRE hazards and TC hazards. HFS linked by key (see Gaps G-06)."),
    ("SOURCES", "Hazard Identification - HAZOP_ODD_v2.xlsx; 2026_SDT_HIRE_from_HAZOP_v1.xlsx; TCs list V3_1.xlsx; Full_Mapping_TC_FI_Hazard_Arch.xlsx; outputs/L1-L5_ODD_report_final.json (map 2026.06.09a); Behavior Table (Jama, M. Pittavino, 2026-03-21); Torc Miro 'Example Lane Change', 'Phase 1A / 1B', 'Concrete Scenario A / B'; Texas Transportation Code 545.053, 545.057, 545.060, 545.104, 545.157."),
    ("STATUS", "All content PROPOSED by SERYTI for review by Torc Safety (Matthieu). No threshold is validated. Colour bands = scenario family. Orange family LC-OOD = [Out-of-ODD] events."),
]
for i, (k, v) in enumerate(lines, 4):
    ws.cell(row=i, column=1, value=k).font = Font(name=FONT, bold=True, size=10, color="1F3864")
    c = ws.cell(row=i, column=2, value=v)
    c.font = Font(name=FONT, size=10)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=i, column=1).alignment = Alignment(vertical="top")
ws.column_dimensions["A"].width = 26
ws.column_dimensions["B"].width = 130

# ---------- Library Tree ----------
ws = out.create_sheet("Library Tree")
title(ws, "Library tree: behavior > family > abstract scenario", "Same layout as a Foretellix V-Suite browsed by family.")
rows = [("Lane Changes / Overtaking (Highway)", "", "", "", "", "")]
r = 4
ws.cell(row=r, column=1, value="Behavior").font = Font(name=FONT, bold=True)
hdr = ["Level", "Family", "ID", "Scenario name (OSC2)", "Torc Miro ref", "Description", "ODD status"]
tree_rows = []
for fid, fname, fdesc in D.FAMILIES:
    tree_rows.append(["Family", fid, fid, fname, "", fdesc, ""])
    for s in [x for x in D.S if x["fam"] == fid]:
        tree_rows.append(["  Scenario", fid, s["id"], s["name"], s["torc"], s["desc"], s["odd"]])
write_table(ws, 4, hdr, tree_rows, [12, 10, 12, 52, 26, 80, 26], fam_col=1)
for i, tr in enumerate(tree_rows, 5):
    if tr[0] == "Family":
        for j in range(1, 8):
            ws.cell(row=i, column=j).font = Font(name=FONT, size=10, bold=True)

# ---------- Scenario Catalogue ----------
ws = out.create_sheet("Scenario Catalogue")
title(ws, "Scenario Catalogue: abstract scenarios for Lane Change (Foretellix style)",
      "One row = one abstract scenario. Parameters stay symbolic here; logical ranges in 'Parameters'. Wrap is on: widen rows to read.")
hdr = ["ID", "Family", "Scenario name (OSC2)", "Torc Miro ref", "Description", "LC type", "Side",
       "Actors (slot: role, action)", "Phases", "Static ODD elements", "Key parameters (symbolic)",
       "Variants", "Wanted behavior (rule book)", "Tasks needed (see Task Chain)", "Coverage items [range, bucket]",
       "KPIs & checkers", "Compatible modifiers", "ODD status", "Hazard contexts stressed",
       "# HIRE rows linked", "Unwanted behaviors (HIRE malfunctions)", "Hazards", "HFS (key)",
       "Triggering conditions (TC-CRAWL)", "Functional insufficiencies (FI)", "Suggested test environment",
       "Public Foretellix analog", "Status"]
rows = []
for s in D.S:
    rows.append([s["id"], s["fam"], s["name"], s["torc"], s["desc"], s["lc_type"], s["side"], s["actors"],
                 s["phases"], s["static"], s["params"], s["variants"], s["wanted"], s["tasks"], s["cov"],
                 s["kpis"], s["mods"], s["odd"], ", ".join(s["ctx"]), len(s["hire_rows"]),
                 ", ".join(s["mal"]), ", ".join(s["hz"]), s["hfs"], s["tc_txt"], s["fi_txt"], s["env"], s["fx"], "PROPOSED"])
widths = [11, 9, 34, 18, 40, 14, 12, 40, 30, 32, 34, 30, 48, 16, 42, 22, 18, 18, 30, 9, 36, 10, 36, 48, 52, 26, 30, 11]
write_table(ws, 4, hdr, rows, widths, fam_col=1)
for i in range(5, 5 + len(rows)):
    ws.row_dimensions[i].height = 150
ws["J4"].comment = Comment("Jama IDs from outputs/L1-L5_ODD_report_final.json (odd_element_definitions).", "SERYTI")
ws["T4"].comment = Comment("Count of Highway Lane Change HIRE rows whose HAZOP Hazard Context is stressed by this scenario. Full list in sheet 'HIRE Traceability'.", "SERYTI")
ws["Y4"].comment = Comment("Derived from the scenario TCs through Full_Mapping_TC_FI_Hazard_Arch.xlsx ('TC-FI exploded').", "SERYTI")

# ---------- Actor Slots ----------
ws = out.create_sheet("Actor Slots")
title(ws, "Actor slot model (from Torc Miro 'Concrete Scenario A / B')",
      "Miro Concrete B: B = Torc lead vehicle, D = Torc chase vehicle, E = controlled vehicle with 2 variants (holds SDT speed / accelerates).")
write_table(ws, 4, ["Slot", "Role", "Position", "Notes"], [list(x) for x in D.SLOTS], [8, 30, 40, 90])
used = {}
for s in D.S:
    for slot in ["A", "B", "C", "D", "E", "F", "G", "S", "O"]:
        if f"{slot}:" in s["actors"] or f"{slot} (" in s["actors"] or f" {slot}," in s["actors"] or f", {slot}" in s["actors"]:
            used.setdefault(slot, []).append(s["id"])

# ---------- Parameters ----------
ws = out.create_sheet("Parameters")
title(ws, "Parameter dictionary: logical ranges for DO-Crawl I-35 South",
      "Status: ODD-quantified = from L1-L5_ODD_report_final.json; proposed = engineering default to review; GAP = needs L4 / ODD data before logical sampling.")
prow = [list(p) for p in D.PARAMS]
write_table(ws, 4, ["Parameter", "Meaning", "Unit", "Logical range (I-35 crawl)", "Status", "Source / note"], prow, [30, 40, 9, 40, 16, 90])
STATUS_FILL = {"ODD-quantified": "C6EFCE", "proposed": "FFEB9C", "GAP": "FFC7CE"}
for i, p in enumerate(prow, 5):
    ws.cell(row=i, column=5).fill = PatternFill("solid", fgColor=STATUS_FILL[p[4]])

# ---------- KPIs ----------
ws = out.create_sheet("KPIs & Checkers")
title(ws, "Scenario evaluation KPIs and checkers",
      "Categories follow the Torc Behavior Table columns for Lane Changes / Overtaking. Thresholds TBD unless noted; candidates must be verified.")
krows = [list(k) for k in D.KPIS]
for k in krows:
    k.append(sum(1 for s in D.S if k[0] in split(s["kpis"])))
write_table(ws, 4, ["KPI ID", "Category", "Metric (OSC2 name)", "Definition", "Unit", "Threshold", "Type", "# scenarios using"], krows, [11, 16, 34, 60, 10, 44, 14, 10])

# ---------- Task Chain ----------
ws = out.create_sheet("Task Chain")
title(ws, "SDT task chain for Lane Change (solution independent)",
      "Answers the Miro sticky 'Don't have Tasks' and the Objectives items 'internal factors along the task chain' and 'KPIs on I/Os of the tasks (solution independent)'. SOTIF team stops here.")
trows = []
for t in D.TASKS:
    n = sum(1 for s in D.S if t[0] in split(s["tasks"].replace(" (fallback planner)", "")))
    trows.append(list(t) + [n])
write_table(ws, 4, ["Task", "Task (what the SDT must do)", "Inputs", "Outputs", "KPIs on task I/O (solution independent)", "Internal factors (FI) along this task", "# scenarios needing it"],
            trows, [7, 48, 32, 36, 60, 48, 11])
for i in range(5, 5 + len(trows)):
    ws.row_dimensions[i].height = 90

# ---------- Modifiers ----------
ws = out.create_sheet("Modifiers")
title(ws, "Modifiers (mixable with any compatible scenario)", "Foretellix 'mix' concept: a modifier changes conditions without changing the scenario logic. Rain, wind, glare are modifiers, not separate scenarios.")
mrows = [list(m) for m in D.MODIFIERS]
for m in mrows:
    m.append(sum(1 for s in D.S if m[0] in split(s["mods"])))
write_table(ws, 4, ["Modifier", "Condition", "ODD element", "Effect on the scenario", "TCs", "Applies to", "# scenarios listing it"], mrows, [9, 32, 34, 50, 44, 18, 11])

# ---------- HIRE Traceability ----------
ws = out.create_sheet("HIRE Traceability")
title(ws, "Phase 1B link: Highway Lane Change HIRE rows > scenarios that exercise them",
      "Source: 2026_SDT_HIRE_from_HAZOP_v1.xlsx, Use Case = Highway, Behavior = Lane Changes / Overtaking (61 rows).")
hrows = [[h["id"], h["odd"], h["man"], h["ctx"], h["gw"], h["mal"], h["ev"], h["hz"], h["S"], h["E"], h["C"], ", ".join(h["sc"]), None] for h in HIRE]
write_table(ws, 4, ["HIRE ID", "ODD condition", "Maneuver", "Hazard context", "Guideword", "Malfunction (unwanted behavior)", "Hazardous event", "Hazard", "S", "E", "C", "Exercised by scenarios", "# scenarios"],
            hrows, [52, 22, 18, 26, 20, 26, 50, 8, 5, 5, 5, 60, 10])
for i in range(5, 5 + len(hrows)):
    ws.cell(row=i, column=13, value=f'=IF(L{i}="",0,LEN(L{i})-LEN(SUBSTITUTE(L{i},",",""))+1)')
    ws.cell(row=i, column=13).font = Font(name=FONT, size=9)
    ws.cell(row=i, column=13).border = BORDER
hire_last = 4 + len(hrows)

# ---------- Scenario x TC x FI ----------
ws = out.create_sheet("Scenario x TC x FI")
title(ws, "Exploded join: scenario x triggering condition x functional insufficiency", "One row per (scenario, TC, FI). Filterable for the FI tree / SOTIF Cl.7 work.")
xrows = []
for s in D.S:
    for t in split(s["tcs"]):
        for f in (TC[t]["fi"] or ["(none mapped)"]):
            xrows.append([s["id"], s["fam"], s["name"], t, TC[t]["layer"], TC[t]["element"], TC[t]["tc"], TC[t]["haz"], f, FI.get(f, {}).get("block", ""), FI.get(f, {}).get("name", "")])
write_table(ws, 4, ["Scenario", "Family", "Scenario name", "TC ID", "Layer", "ODD element", "Triggering condition", "TC hazards", "FI ID", "FI block", "Functional insufficiency"], xrows, [11, 9, 40, 14, 6, 26, 50, 12, 11, 16, 60], fam_col=1)
x_last = 4 + len(xrows)

# ---------- Coverage Stats (formulas) ----------
ws = out.create_sheet("Coverage Stats")
title(ws, "Coverage stats (live formulas)")
cat_last = 4 + len(D.S)
r = 4
ws.cell(row=r, column=1, value="Family"); ws.cell(row=r, column=2, value="Name"); ws.cell(row=r, column=3, value="# scenarios"); ws.cell(row=r, column=4, value="# HIRE links (sum)")
style_header(ws, r, 4)
for i, (fid, fname, _) in enumerate(D.FAMILIES, r + 1):
    ws.cell(row=i, column=1, value=fid); ws.cell(row=i, column=2, value=fname)
    ws.cell(row=i, column=3, value=f"=COUNTIF('Scenario Catalogue'!$B$5:$B${cat_last},A{i})")
    ws.cell(row=i, column=4, value=f"=SUMIF('Scenario Catalogue'!$B$5:$B${cat_last},A{i},'Scenario Catalogue'!$T$5:$T${cat_last})")
tot = r + len(D.FAMILIES) + 1
ws.cell(row=tot, column=2, value="Total")
ws.cell(row=tot, column=3, value=f"=SUM(C{r+1}:C{tot-1})")
ws.cell(row=tot, column=4, value=f"=SUM(D{r+1}:D{tot-1})")
r = tot + 2
ws.cell(row=r, column=1, value="Check"); ws.cell(row=r, column=2, value="Value"); ws.cell(row=r, column=3, value="Expected")
style_header(ws, r, 3)
checks = [
    ("HIRE Lane Change rows", f"=COUNTA('HIRE Traceability'!A5:A{hire_last})", "61"),
    ("HIRE rows exercised by >= 1 scenario", f"=COUNTIF('HIRE Traceability'!M5:M{hire_last},\">0\")", "equal to total"),
    ("HIRE rows not exercised", f"=COUNTIF('HIRE Traceability'!M5:M{hire_last},0)", "0"),
    ("Scenario x TC x FI rows", f"=COUNTA('Scenario x TC x FI'!A5:A{x_last})", ""),
    ("Distinct TCs used (see note)", len({t for s in D.S for t in split(s["tcs"])}), "computed at build"),
    ("Distinct FIs reached (see note)", len({f for s in D.S for f in s["fis"]}), "of 42 in FI list"),
    ("Scenarios In-ODD", f"=COUNTIF('Scenario Catalogue'!R5:R{cat_last},\"In-ODD*\")", ""),
    ("Scenarios with [Out-of-ODD] content", f"=COUNTIF('Scenario Catalogue'!R5:R{cat_last},\"*Out-of-ODD*\")", ""),
    ("Parameters marked GAP", f"=COUNTIF(Parameters!E5:E{4+len(D.PARAMS)},\"GAP\")", ""),
]
for i, (k, v, e) in enumerate(checks, r + 1):
    ws.cell(row=i, column=1, value=k); ws.cell(row=i, column=2, value=v); ws.cell(row=i, column=3, value=e)
for row in ws.iter_rows(min_row=4, max_row=r + len(checks)):
    for c in row:
        if c.value is not None and c.row not in (4, r):
            c.font = Font(name=FONT, size=10)
ws.cell(row=r + len(checks) + 2, column=1, value="Note: distinct TC / FI counts are written as values computed by the build script (set operations), not formulas.").font = Font(name=FONT, italic=True, size=9)
ws.column_dimensions["A"].width = 40; ws.column_dimensions["B"].width = 44; ws.column_dimensions["C"].width = 16; ws.column_dimensions["D"].width = 18

# ---------- Gaps ----------
ws = out.create_sheet("Gaps & Open Points")
title(ws, "Gaps and open points")
write_table(ws, 4, ["ID", "Topic", "Gap", "Proposed action / owner"], [list(g) for g in D.GAPS], [7, 20, 90, 70])

for w in out.worksheets:
    w.sheet_view.zoomScale = 90
out.save(OUT)
print("saved", OUT, "HIRE", len(HIRE), "xrows", len(xrows))
print("uncovered HIRE:", [h["id"] for h in HIRE if not h["sc"]])
print("distinct TCs", len({t for s in D.S for t in split(s["tcs"])}), "FIs", len({f for s in D.S for f in s["fis"]}))
