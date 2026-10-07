from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
import json, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

cat = json.load(open(ROOT / "catalogue/catalog_full.json"))
counts = json.load(open(ROOT / "catalogue/counts.json"))
B = cat["behaviors"]; ABS = cat["abstracts"]; FAM = cat["families"]; NH = cat["nhtsa"]; JR = cat["jama_reqs"]; FX = cat["foretellix"]
code = {b["id"]: b["code"] for b in B}
LET = {"keep_lane": "K", "lane_change": "L", "move_over": "M", "stop_on_shoulder": "S", "stop_in_lane": "I", "emergency_lateral": "E"}
F = "Arial"
HDR = PatternFill("solid", fgColor="0A6B46"); FAMF = PatternFill("solid", fgColor="E3EEE8"); ON = PatternFill("solid", fgColor="CFE8DA"); NG = PatternFill("solid", fgColor="FBEBCB")
thin = Side(style="thin", color="C8D2CC"); BD = Border(left=thin, right=thin, top=thin, bottom=thin)
def hdr(ws, row, n):
    for c in range(1, n + 1):
        x = ws.cell(row=row, column=c); x.font = Font(name=F, bold=True, color="FFFFFF", size=10); x.fill = HDR; x.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center"); x.border = BD
def body(ws, r0, r1, n, wrap=True):
    for r in range(r0, r1 + 1):
        for c in range(1, n + 1):
            x = ws.cell(row=r, column=c); x.font = Font(name=F, size=9, bold=x.font.bold); x.alignment = Alignment(wrap_text=wrap, vertical="top"); x.border = BD
def title(ws, t, sub):
    ws["A1"] = t; ws["A1"].font = Font(name=F, bold=True, size=14, color="0A6B46")
    ws["A2"] = sub; ws["A2"].font = Font(name=F, italic=True, size=9, color="56655E")
def actors_txt(a):
    if not a.get("configuration"): return ""
    return "\n".join(f'{x["slot"]} ({ {"required": "obligatoire", "optional": "optionnel", "off": "absent"}[x["mode"]] }) : {", ".join(x["kinds"])}' + "".join(f' · {m["type"]}' for m in x.get("maneuvers", []) if m["type"] != "none") for x in a["configuration"]["actors"]) or "aucun (SDT seul)"

wb = openpyxl.Workbook()
ws = wb.active; ws.title = "Lisez-moi"
title(ws, "Catalogue des scénarios abstraits et lien Behavior × Catalogue (v3, PROPOSED)", "SERYTI pour Torc Robotics, DO-Crawl I-35 South, 7 octobre 2026. Même contenu que l'outil Scenario Studio.")
rows = [
 ("Ce que c'est", f"{len(ABS)} scénarios abstraits (situations uniques) et leurs liens vers les {len(B)} behaviors de la Behavior Table Torc. Un abstract peut vérifier plusieurs behaviors (lien n:m) : {sum(len(a['behaviors']) for a in ABS)} liens au total."),
 ("D'où viennent les abstracts", "Écrits par SERYTI. Ils ne viennent pas de la bibliothèque Foretellix (sous licence) : seule la structure (acteurs, phases, paramètres, coverage) suit le format Foretellix. Sources : Miro Torc (Concrete A/B), HAZOP/HIRE v2, catalogue TC-CRAWL V3, typologie NHTSA pré-crash (DOT HS 810 767), définitions Jama des behaviors (export du 26/03/2026)."),
 ("Foretellix", "La bibliothèque publique Foretellix / SAFE « Highway and ADAS Traffic Scenario Library » (VMAD-SG1-11-06, novembre 2020, 13 scénarios fonctionnels) a été croisée avec le catalogue : 9 abstracts ont été ajoutés pour les situations applicables qui manquaient (convoi lent à côté, lane hugger, file arrêtée à côté, courbe avec voisin, leader qui sort, convoi qui s'insère, entrecroisement à l'insertion, double cut-in, cut-out simple). Onglet « Couverture Foretellix ». Le catalogue Torc réel est la « L4 Highway Trucking V-Suite » de Foretellix, non publique : à demander (scenario index and coverage model)."),
 ("Justification HIRE", "Chaque lien abstract → behavior est justifié par des lignes du HIRE v2 (2026_SDT_HIRE_from_HAZOP_v2, 1 156 HARA) et rattaché aux HFS (Hazardous_Functional_Scenarios_V1). Règle : « justifié » quand le HIRE du behavior relie une commande sollicitée par la scène (décélération, braquage, accélération, clignotant, feux) au hazard pour lequel l'abstract a été écrit ; « justifié indirectement » quand c'est un autre hazard exposé par la scène, ou quand l'abstract n'est pas encore générable (hazard seul). Onglets « Justification HIRE » (une ligne par lien) et « Trous HIRE » (combinaisons du HIRE sans abstract)."),
 ("Ce que l'analyse a changé", "5 liens ajoutés parce que le HIRE contient la condition ODD pour ce behavior et qu'aucun abstract ne la couvrait (OBJ-04, LC-OOD-04 et OBJ-06 pour Cruise ; RD-13 et OBJ-02 pour Controlled Stop on Shoulder), et 1 abstract créé : RMP-14 (bretelle d'accès en courbe et en pente, H5 de On-Ramp Merge)."),
 ("Ce qui manque", "Le catalogue de Matthieu (environ 104 abstracts, colonne A famille, colonne B abstract, nombre à confirmer) n'est pas dans le dossier. La colonne « Équivalent catalogue Matthieu » de l'onglet Catalogue est prévue pour faire le rapprochement."),
 ("Intention du SDT", "Chaque behavior impose une intention au SDT pour la génération : K rester dans la voie, L changement de voie, M se décaler, S arrêt sur l'accotement, I arrêt en voie, E évitement latéral. L'intention est observée sur l'ADS, jamais scriptée dans l'OSC."),
 ("Génération", "Les abstracts autoroute et les abstracts urbains en voie sont générables dans l'outil (acteurs A à O). Les intersections, la marche arrière, le contresens et la cour du hub ne le sont pas encore (jalon 2) : ils sont dans le catalogue et dans la matrice, sans génération."),
 ("Exigences Jama", "Liens proposés vers les exigences de l'export « Driver Out 2026 SDT Requirements » (définitions des behaviors, Lane Change, Stay in Lane, OMV, vitesses) et vers TORC-SYSRQ-15232 et TORC-FSR-3752 (instantané local du repo Scenario Studio). À relire."),
 ("Onglets", "Catalogue · Matrice Behavior × Abstract · Justification HIRE · Trous HIRE · Behaviors · Couverture NHTSA · Couverture Foretellix · Exigences Jama."),
 ("Statut", "PROPOSED. Aucun lien n'est validé par Torc."),
]
for i, (k, v) in enumerate(rows, 4):
    ws.cell(row=i, column=1, value=k).font = Font(name=F, bold=True, size=10, color="0A6B46")
    c = ws.cell(row=i, column=2, value=v); c.font = Font(name=F, size=10); c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=i, column=1).alignment = Alignment(vertical="top")
ws.column_dimensions["A"].width = 26; ws.column_dimensions["B"].width = 130

# Catalogue
ws = wb.create_sheet("Catalogue")
title(ws, "Catalogue : une ligne par scénario abstrait", "Concrets = combinaisons générées par l'outil avec les valeurs par défaut, sommées sur les behaviors liés.")
H = ["ID", "Famille", "Titre", "Description", "Use cases", "Behaviors liés", "Nb behaviors", "Acteurs (slot, présence, types, manœuvres)", "Générable", "Combinaisons", "Concrets retenus", "Concrets exclus", "NHTSA pré-crash", "Hazards", "Triggering conditions", "Exigences Jama", "Dépend de la perception", "Source", "Statut", "Équivalent catalogue Matthieu", "Foretellix SAFE (VMAD 2020)"]
for j, h in enumerate(H, 1): ws.cell(row=4, column=j, value=h)
hdr(ws, 4, len(H))
r = 5
for a in ABS:
    cs = counts.get(a["id"], {})
    vals = [a["id"], a["family"], a.get("title") or a["name"], a["description"], ", ".join(a.get("use_cases", [])), ", ".join(code[b] for b in a["behaviors"]), None,
            actors_txt(a), "oui" if a.get("configuration") else "jalon 2", sum(x["potential"] for x in cs.values()) if cs else None, sum(x["valid"] for x in cs.values()) if cs else None, sum(x["rejected"] for x in cs.values()) if cs else None,
            ", ".join(f'#{n} {NH[str(n)]}' for n in a.get("nhtsa", [])), ", ".join(a.get("hazards", [])), ", ".join(a.get("tcs", [])), ", ".join(a.get("jama", [])),
            "oui" if a.get("perception_dependent") else "non", a.get("source", ""), a.get("status", ""), "", ", ".join(f'§{r["ref"]} {FX[r["ref"]]["title"]} ({r["level"]})' for r in a.get("foretellix", []))]
    for j, v in enumerate(vals, 1): ws.cell(row=r, column=j, value=v)
    ws.cell(row=r, column=7, value=f'=IF(F{r}="",0,LEN(F{r})-LEN(SUBSTITUTE(F{r},",",""))+1)')
    r += 1
last_cat = r - 1
body(ws, 5, last_cat, len(H))
for j, w in enumerate([11, 8, 34, 50, 12, 26, 9, 40, 10, 11, 11, 10, 34, 10, 30, 36, 11, 22, 11, 26, 30], 1): ws.column_dimensions[get_column_letter(j)].width = w
ws.freeze_panes = "D5"; ws.auto_filter.ref = f"A4:{get_column_letter(len(H))}{last_cat}"
tr = last_cat + 2
ws.cell(row=tr, column=3, value="Totaux").font = Font(name=F, bold=True)
for col in "JKL": ws[f"{col}{tr}"] = f"=SUM({col}5:{col}{last_cat})"; ws[f"{col}{tr}"].font = Font(name=F, bold=True)
ws[f"G{tr}"] = f"=SUM(G5:G{last_cat})"; ws[f"G{tr}"].font = Font(name=F, bold=True)

# Matrix
ws = wb.create_sheet("Matrice")
title(ws, "Matrice Behavior × Abstract", "Lettre = intention du SDT pour ce behavior (K, L, M, S, I, E). « 2 » = lien prévu, génération au jalon 2.")
ws.cell(row=4, column=1, value="ID"); ws.cell(row=4, column=2, value="Abstract")
for j, b in enumerate(B, 3): ws.cell(row=4, column=j, value=b["code"])
nb = len(B); ws.cell(row=4, column=nb + 3, value="Nb behaviors")
hdr(ws, 4, nb + 3)
for j in range(3, nb + 3): ws.cell(row=4, column=j).alignment = Alignment(text_rotation=90, horizontal="center", vertical="bottom")
ws.row_dimensions[4].height = 70
r = 5; fam = None
for a in ABS:
    if a["family"] != fam:
        fam = a["family"]; ws.cell(row=r, column=1, value=f'{fam} · {FAM.get(fam, {}).get("name", "")}')
        for j in range(1, nb + 4): ws.cell(row=r, column=j).fill = FAMF
        ws.cell(row=r, column=1).font = Font(name=F, bold=True, size=9); r += 1
    ws.cell(row=r, column=1, value=a["id"]); ws.cell(row=r, column=2, value=a.get("title") or a["name"])
    for j, b in enumerate(B, 3):
        if b["id"] in a["behaviors"]:
            if a.get("configuration"):
                v = LET.get(counts[a["id"]][b["code"]]["intent"], "x"); fill = ON
            else:
                v = "2"; fill = NG
            c = ws.cell(row=r, column=j, value=v); c.fill = fill; c.alignment = Alignment(horizontal="center")
    ws.cell(row=r, column=nb + 3, value=f'=COUNTA(C{r}:{get_column_letter(nb + 2)}{r})')
    r += 1
last_mx = r - 1
body(ws, 5, last_mx, nb + 3, wrap=False)
for rr in range(5, last_mx + 1):
    for j in range(3, nb + 3): ws.cell(row=rr, column=j).alignment = Alignment(horizontal="center")
cr = last_mx + 1
ws.cell(row=cr, column=2, value="Abstracts par behavior").font = Font(name=F, bold=True)
for j in range(3, nb + 3):
    col = get_column_letter(j); ws.cell(row=cr, column=j, value=f'=COUNTA({col}5:{col}{last_mx})').font = Font(name=F, bold=True)
ws.column_dimensions["A"].width = 12; ws.column_dimensions["B"].width = 48
for j in range(3, nb + 3): ws.column_dimensions[get_column_letter(j)].width = 5.5
ws.column_dimensions[get_column_letter(nb + 3)].width = 10
ws.freeze_panes = "C5"

# Behaviors
ws = wb.create_sheet("Behaviors")
title(ws, "Behaviors (Behavior Table Jama) et leur catalogue", "Définition Jama : export « Driver Out 2026 SDT Requirements » du 26/03/2026 (Highway seulement).")
H = ["Code", "Use case", "Behavior", "DDT / DDT-F", "Catégories", "Intention du SDT", "Définition Jama", "Statut Jama", "Nb abstracts", "Concrets retenus", "Texte de la définition"]
for j, h in enumerate(H, 1): ws.cell(row=4, column=j, value=h)
hdr(ws, 4, len(H)); r = 5
for b in B:
    tot = sum(counts[a["id"]].get(b["code"], {}).get("valid", 0) for a in ABS if a.get("configuration") and b["id"] in a["behaviors"])
    vals = [b["code"], b["use_case"], b["name"], b["ddt"], ", ".join(b.get("categories", [])), b["default"]["intent"], (b.get("jama") or {}).get("id", ""), (b.get("jama") or {}).get("status", ""), None, tot, (b.get("jama") or {}).get("text", "")]
    for j, v in enumerate(vals, 1): ws.cell(row=r, column=j, value=v)
    col = get_column_letter(3 + B.index(b))
    ws.cell(row=r, column=9, value=f"=Matrice!{col}{cr}")
    r += 1
body(ws, 5, r - 1, len(H))
for j, w in enumerate([10, 10, 34, 10, 30, 16, 18, 10, 10, 12, 100], 1): ws.column_dimensions[get_column_letter(j)].width = w
ws.freeze_panes = "D5"

# Justification HIRE
HZN = cat["hazard_names"]; H2 = cat["hire2"]
ws = wb.create_sheet("Justification HIRE")
title(ws, "Justification de chaque lien abstract → behavior par le HIRE v2", "Une ligne par lien. Lignes HIRE : les premières sont celles de la condition ODD de l'abstract, puis les conditions nominales.")
H = ["Abstract", "Titre", "Behavior", "Niveau", "Origine du lien", "Pourquoi", "Hazards", "Commandes sollicitées", "Nb lignes HIRE", "Lignes HIRE v2 (10 premières)", "HFS", "Condition ODD dans le HIRE", "Condition ODD absente du HIRE"]
for j, h in enumerate(H, 1): ws.cell(row=4, column=j, value=h)
hdr(ws, 4, len(H)); r = 5
LV = {"justifié": ON, "justifié indirectement": NG}
for a in ABS:
    for c in a["beh_codes"]:
        L = a["hire_links"][c]
        vals = [a["id"], a.get("title") or a["name"], c, L["level"], (a.get("link_origin") or {}).get(c, "catalogue SERYTI"), L["why"], ", ".join(f"{z} {HZN.get(z, '')}" for z in L["hz"]), ", ".join(L["cmd"]), L["n"], "\n".join(L["rows"][:10]), ", ".join(L["hfs"]), ", ".join(L["odd_match"]), ", ".join(L["odd_missing"])]
        for j, v in enumerate(vals, 1): ws.cell(row=r, column=j, value=v)
        r += 1
last = r - 1
body(ws, 5, last, len(H))
for rr in range(5, last + 1):
    f = LV.get(ws.cell(row=rr, column=4).value)
    if f: ws.cell(row=rr, column=4).fill = f
for j, w in enumerate([11, 32, 9, 15, 26, 48, 26, 30, 9, 44, 18, 30, 30], 1): ws.column_dimensions[get_column_letter(j)].width = w
ws.freeze_panes = "D5"; ws.auto_filter.ref = f"A4:{get_column_letter(len(H))}{last}"
ws.cell(row=last + 2, column=2, value="Liens").font = Font(name=F, bold=True); ws.cell(row=last + 2, column=4, value=f"=COUNTA(C5:C{last})").font = Font(name=F, bold=True)
ws.cell(row=last + 3, column=2, value="Justifiés").font = Font(name=F, bold=True); ws.cell(row=last + 3, column=4, value=f'=COUNTIF(D5:D{last},"justifié")').font = Font(name=F, bold=True)
ws.cell(row=last + 4, column=2, value="Justifiés indirectement").font = Font(name=F, bold=True); ws.cell(row=last + 4, column=4, value=f'=COUNTIF(D5:D{last},"justifié indirectement")').font = Font(name=F, bold=True)
ws.cell(row=last + 5, column=2, value="À relire").font = Font(name=F, bold=True); ws.cell(row=last + 5, column=4, value=f'=COUNTIF(D5:D{last},"à relire")').font = Font(name=F, bold=True)

# Trous HIRE
ws = wb.create_sheet("Trous HIRE")
title(ws, "Trous du catalogue vus depuis le HIRE v2", "Combinaisons du HIRE (behavior × hazard, behavior × condition ODD) et abstracts qui les couvrent. Les états système « All / State » (84 lignes) ne sont pas des behaviors et ne sont pas comptés.")
H = ["Use case", "Behavior", "Type", "Combinaison HIRE", "Lignes HIRE", "dont MRM", "Malfunctions", "Abstracts qui couvrent", "Nb abstracts", "Statut", "Action proposée"]
for j, h in enumerate(H, 1): ws.cell(row=4, column=j, value=h)
hdr(ws, 4, len(H)); r = 5
for g in sorted(cat["hire_gaps"], key=lambda g: (["Highway", "Urban", "Hub"].index(g["uc"]), g["code"], g["kind"], g["key"])):
    act = ("non modélisé au jalon 1 (perception parfaite, adhérence)" if g.get("not_modelled") else ("lier " + ", ".join(g["candidates"]) if g.get("candidates") else ("nouvel abstract" if g["code"].startswith("HW") else "catalogue urbain / hub à compléter (jalon 2)"))) if not g["covered_by"] else ""
    vals = [g["uc"], g["code"], "Hazard" if g["kind"] == "hazard" else "Condition ODD", g["label"], g["n"], g["mrm"] or None, ", ".join(g["mals"][:8]), ", ".join(g["covered_by"]), None, None, act]
    for j, v in enumerate(vals, 1): ws.cell(row=r, column=j, value=v)
    ws.cell(row=r, column=9, value=f'=IF(H{r}="",0,LEN(H{r})-LEN(SUBSTITUTE(H{r},",",""))+1)')
    ws.cell(row=r, column=10, value=f'=IF(I{r}>0,"couvert","trou")')
    r += 1
last = r - 1
body(ws, 5, last, len(H))
for j, w in enumerate([10, 10, 13, 40, 9, 8, 50, 50, 9, 9, 42], 1): ws.column_dimensions[get_column_letter(j)].width = w
ws.freeze_panes = "D5"; ws.auto_filter.ref = f"A4:{get_column_letter(len(H))}{last}"
for i, (lab, f) in enumerate([("Combinaisons", f"=COUNTA(B5:B{last})"), ("Couvertes", f'=COUNTIF(J5:J{last},"couvert")'), ("Trous", f'=COUNTIF(J5:J{last},"trou")'), ("Trous Highway", f'=COUNTIFS(A5:A{last},"Highway",J5:J{last},"trou")')]):
    ws.cell(row=last + 2 + i, column=4, value=lab).font = Font(name=F, bold=True); ws.cell(row=last + 2 + i, column=5, value=f).font = Font(name=F, bold=True)

# NHTSA
ws = wb.create_sheet("Couverture NHTSA")
title(ws, "Couverture de la typologie NHTSA pré-crash (DOT HS 810 767)", "37 typologies de pré-crash ; abstracts qui en couvrent chacune.")
H = ["#", "Typologie NHTSA", "Abstracts", "Nb abstracts", "Couverte"]
for j, h in enumerate(H, 1): ws.cell(row=4, column=j, value=h)
hdr(ws, 4, len(H)); r = 5
for n in range(1, 38):
    lst = [a["id"] for a in ABS if n in a.get("nhtsa", [])]
    ws.cell(row=r, column=1, value=n); ws.cell(row=r, column=2, value=NH[str(n)]); ws.cell(row=r, column=3, value=", ".join(lst))
    ws.cell(row=r, column=4, value=f'=IF(C{r}="",0,LEN(C{r})-LEN(SUBSTITUTE(C{r},",",""))+1)'); ws.cell(row=r, column=5, value=f'=IF(D{r}>0,"oui","non")')
    r += 1
body(ws, 5, r - 1, len(H))
ws.cell(row=r + 1, column=2, value="Typologies couvertes").font = Font(name=F, bold=True)
ws.cell(row=r + 1, column=4, value=f'=COUNTIF(E5:E{r - 1},"oui")').font = Font(name=F, bold=True)
ws.cell(row=r + 2, column=2, value="Note : #1 (pas de conducteur), #21 (manœuvre en sens inverse) et #35 (incident sans collision) ne s'appliquent pas ou peu à un camion autonome sur I-35.").font = Font(name=F, italic=True, size=9)
for j, w in enumerate([5, 48, 70, 11, 10], 1): ws.column_dimensions[get_column_letter(j)].width = w

# Foretellix
ws = wb.create_sheet("Couverture Foretellix")
title(ws, "Couverture de la bibliothèque Foretellix / SAFE (VMAD-SG1-11-06, novembre 2020)", "13 scénarios fonctionnels publics (le §4 en contient 3). Couvert = l'abstract reproduit la situation ; partiel = une partie seulement.")
H = ["§", "Scénario Foretellix", "Description", "Applicable à Torc (camion L4, I-35)", "Abstracts qui couvrent", "Abstracts partiels", "Nb couvrants", "Statut"]
for j, h in enumerate(H, 1): ws.cell(row=4, column=j, value=h)
hdr(ws, 4, len(H)); r = 5
for ref, fx in FX.items():
    cov = [a["id"] for a in ABS if any(x["ref"] == ref and x["level"] == "couvert" for x in a.get("foretellix", []))]
    par = [a["id"] for a in ABS if any(x["ref"] == ref and x["level"] == "partiel" for x in a.get("foretellix", []))]
    for j, v in enumerate([ref, fx["title"], fx["desc"], fx["torc"], ", ".join(cov), ", ".join(par)], 1): ws.cell(row=r, column=j, value=v)
    ws.cell(row=r, column=7, value=f'=IF(E{r}="",0,LEN(E{r})-LEN(SUBSTITUTE(E{r},",",""))+1)')
    ws.cell(row=r, column=8, value=f'=IF(G{r}>0,"couvert",IF(F{r}<>"","partiel","non couvert"))')
    r += 1
body(ws, 5, r - 1, len(H))
ws.cell(row=r + 1, column=2, value="Scénarios couverts").font = Font(name=F, bold=True)
ws.cell(row=r + 1, column=7, value=f'=COUNTIF(H5:H{r - 1},"couvert")').font = Font(name=F, bold=True)
ws.cell(row=r + 2, column=2, value="Ajoutés en v2 grâce à cette bibliothèque : ADJ-05, ADJ-06, ADJ-07, RD-14, RMP-11, RMP-12, RMP-13, CI-08, LV-11.").font = Font(name=F, italic=True, size=9)
for j, w in enumerate([6, 26, 50, 40, 44, 34, 11, 12], 1): ws.column_dimensions[get_column_letter(j)].width = w
ws.freeze_panes = "C5"

# Jama
ws = wb.create_sheet("Exigences Jama")
title(ws, "Exigences Jama et abstracts qui permettent de les vérifier", "Liens proposés par SERYTI, à relire.")
H = ["ID", "Nom", "Statut", "Abstracts liés", "Nb abstracts", "Texte"]
for j, h in enumerate(H, 1): ws.cell(row=4, column=j, value=h)
hdr(ws, 4, len(H)); r = 5
for rid, rq in sorted(JR.items()):
    lst = [a["id"] for a in ABS if rid in a.get("jama", [])]
    for j, v in enumerate([rid, rq["name"], rq.get("status"), ", ".join(lst), None, rq["text"]], 1): ws.cell(row=r, column=j, value=v)
    ws.cell(row=r, column=5, value=f'=IF(D{r}="",0,LEN(D{r})-LEN(SUBSTITUTE(D{r},",",""))+1)')
    r += 1
body(ws, 5, r - 1, len(H))
for j, w in enumerate([19, 46, 10, 60, 10, 90], 1): ws.column_dimensions[get_column_letter(j)].width = w
ws.freeze_panes = "C5"
for w in wb.worksheets: w.sheet_view.zoomScale = 90
wb.save(ROOT / "catalogue/Catalogue_Scenarios_Abstraits_v3.xlsx")
print("saved")
